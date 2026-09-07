from openpyxl.styles import (
    Alignment,
    PatternFill,
    Font,
    Border,
    Side
)


# -----------------------------------
# Fills
# -----------------------------------

def get_fills():

    blue_fill = PatternFill(
        fill_type="solid",
        fgColor="4472C4"
    )

    gray_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E1F2"
    )

    beige_fill = PatternFill(
        fill_type="solid",
        fgColor="F5E6CC"
    )

    return {
        "blue": blue_fill,
        "gray": gray_fill,
        "beige": beige_fill
    }


# -----------------------------------
# Fonts
# -----------------------------------

def get_fonts():

    white_font = Font(
        color="FFFFFF",
        bold=True
    )

    bold_font = Font(
        bold=True
    )

    summary_header_font = Font(
        color="FFFFFF",
        bold=True,
        size=12
    )

    return {
        "white": white_font,
        "bold": bold_font,
        "summary_header": summary_header_font
    }


# -----------------------------------
# Borders
# -----------------------------------

def get_borders():

    thin_side = Side(
        style="thin",
        color="000000"
    )

    medium_side = Side(
        style="medium",
        color="000000"
    )

    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    section_border = Border(
        top=medium_side,
        bottom=medium_side
    )

    return {
        "thin": thin_border,
        "section": section_border
    }


# -----------------------------------
# Common alignments
# -----------------------------------

def center_alignment():

    return Alignment(
        horizontal="center",
        vertical="center"
    )


def left_alignment():

    return Alignment(
        horizontal="left",
        vertical="center"
    )