from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from styles import (
    get_fills,
    get_fonts,
    center_alignment
)


# -----------------------------------
# Create workbook
# -----------------------------------

def create_workbook():

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Weight Data"

    return workbook, sheet


# -----------------------------------
# Create main header
# -----------------------------------

def create_main_header(
    sheet,
    total_columns
):

    fills = get_fills()
    fonts = get_fonts()

    # -----------------------------------
    # First header
    # -----------------------------------

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=total_columns
    )

    sheet["A1"] = (
        "HPS INTEGRATORS CONTRACT ANIMAL GROWING SERVICE"
    )

    sheet["A1"].fill = fills["blue"]

    sheet["A1"].font = fonts["white"]

    sheet["A1"].alignment = center_alignment()


    # -----------------------------------
    # Second header
    # -----------------------------------

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=total_columns
    )

    sheet["A2"] = "LIVE SALES TALLY SHEET"

    sheet["A2"].fill = fills["blue"]

    sheet["A2"].font = fonts["white"]

    sheet["A2"].alignment = center_alignment()


    # -----------------------------------
    # Header row heights
    # -----------------------------------

    sheet.row_dimensions[1].height = 25

    sheet.row_dimensions[2].height = 22


# -----------------------------------
# Create information section
# -----------------------------------

def create_information_section(
    sheet,
    total_columns,
    time_started,
    file_date
):

    fonts = get_fonts()

    # -----------------------------------
    # Left side
    # -----------------------------------

    sheet["A4"] = "GROWER:"

    sheet["A5"] = "LOCATION:"

    sheet["A6"] = "BUILDING NO:"

    sheet["A7"] = "TIME STARTED:"

    sheet["B7"] = time_started


    # -----------------------------------
    # Bold left labels
    # -----------------------------------

    sheet["A4"].font = fonts["bold"]
    sheet["A5"].font = fonts["bold"]
    sheet["A6"].font = fonts["bold"]
    sheet["A7"].font = fonts["bold"]
    sheet["B7"].font = fonts["bold"]


    # -----------------------------------
    # Right side
    # -----------------------------------

    right_label_column = total_columns - 1

    right_value_column = total_columns


    sheet.cell(
        row=4,
        column=right_label_column
    ).value = "DATE:"


    sheet.cell(
        row=4,
        column=right_value_column
    ).value = file_date


    sheet.cell(
        row=5,
        column=right_label_column
    ).value = "BUYER:"


    sheet.cell(
        row=6,
        column=right_label_column
    ).value = "D.R NO:"


    sheet.cell(
        row=7,
        column=right_label_column
    ).value = "AGE:"


    sheet.cell(
        row=8,
        column=right_label_column
    ).value = "TIME FINISHED:"


    # -----------------------------------
    # Bold right labels
    # -----------------------------------

    for row in range(4, 9):

        sheet.cell(
            row=row,
            column=right_label_column
        ).font = fonts["bold"]


# -----------------------------------
# Create measurement table
# -----------------------------------

def create_measurement_table(
    sheet,
    groups,
    total_columns
):

    fills = get_fills()
    fonts = get_fonts()

    header_row = 10

    data_start_row = 11

    data_end_row = 20

    total_row = 21


    # -----------------------------------
    # Groups
    # -----------------------------------

    for group_number, group in enumerate(groups):

        time_column = (
            group_number * 2 + 1
        )

        weight_column = (
            group_number * 2 + 2
        )


        # -----------------------------------
        # Table headers
        # -----------------------------------

        time_header = sheet.cell(
            row=header_row,
            column=time_column
        )

        weight_header = sheet.cell(
            row=header_row,
            column=weight_column
        )


        time_header.value = "TIME"

        weight_header.value = "20 HDS"


        time_header.fill = fills["gray"]

        weight_header.fill = fills["gray"]


        time_header.font = fonts["bold"]

        weight_header.font = fonts["bold"]


        time_header.alignment = center_alignment()

        weight_header.alignment = center_alignment()


        # -----------------------------------
        # Measurements
        # -----------------------------------

        for row_number, entry in enumerate(
            group,
            start=data_start_row
        ):

            time_cell = sheet.cell(
                row=row_number,
                column=time_column
            )

            weight_cell = sheet.cell(
                row=row_number,
                column=weight_column
            )


            time_cell.value = entry["time"]

            weight_cell.value = entry["weight"]


            # -----------------------------------
            # Beige background
            # -----------------------------------

            time_cell.fill = fills["beige"]

            weight_cell.fill = fills["beige"]


            # -----------------------------------
            # Center
            # -----------------------------------

            time_cell.alignment = center_alignment()

            weight_cell.alignment = center_alignment()


        # -----------------------------------
        # Total
        # -----------------------------------

        total_time_cell = sheet.cell(
            row=total_row,
            column=time_column
        )

        total_weight_cell = sheet.cell(
            row=total_row,
            column=weight_column
        )


        total_time_cell.value = "TOTAL"


        total_weight_cell.value = (
            f"=SUM("
            f"{sheet.cell(row=data_start_row, column=weight_column).coordinate}:"
            f"{sheet.cell(row=data_end_row, column=weight_column).coordinate}"
            f")"
        )


        total_time_cell.font = fonts["bold"]

        total_weight_cell.font = fonts["bold"]


    # -----------------------------------
    # Column widths
    # -----------------------------------

    for column in range(
        1,
        total_columns + 1
    ):

        column_letter = get_column_letter(column)

        if column % 2 == 1:

            # TIME columns
            sheet.column_dimensions[
                column_letter
            ].width = 18

        else:

            # 20 HDS columns
            sheet.column_dimensions[
                column_letter
            ].width = 8



