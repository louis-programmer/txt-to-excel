from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from styles import (
    get_fills,
    get_fonts,
    get_borders,
    center_alignment,
    left_alignment
)


# -----------------------------------
# Create final summary
# -----------------------------------

def create_summary(
    sheet,
    groups,
    total_columns,
    data_start_row=11,
    data_end_row=20,
    total_row=21
):

    fills = get_fills()
    fonts = get_fonts()
    borders = get_borders()

    # -----------------------------------
    # Start final section on new page
    # -----------------------------------

    sheet.row_breaks.append(
        Break(id=total_row + 2)
    )

    final_start_row = total_row + 3


    # -----------------------------------
    # FINAL SUMMARY HEADER
    # -----------------------------------

    summary_header_row = final_start_row

    sheet.merge_cells(
        start_row=summary_header_row,
        start_column=1,
        end_row=summary_header_row,
        end_column=total_columns
    )

    summary_header = sheet.cell(
        row=summary_header_row,
        column=1
    )

    summary_header.value = "FINAL REPORT / SUMMARY"

    summary_header.fill = fills["blue"]

    summary_header.font = fonts["summary_header"]

    summary_header.alignment = center_alignment()

    sheet.row_dimensions[
        summary_header_row
    ].height = 25


    # -----------------------------------
    # REMARKS + TOTALS
    # -----------------------------------

    remarks_row = summary_header_row + 2


    # -----------------------------------
    # REMARKS HEADER
    # -----------------------------------

    remarks_start_column = 1

    remarks_end_column = total_columns // 2

    sheet.merge_cells(
        start_row=remarks_row,
        start_column=remarks_start_column,
        end_row=remarks_row,
        end_column=remarks_end_column
    )

    remarks_header = sheet.cell(
        row=remarks_row,
        column=remarks_start_column
    )

    remarks_header.value = "REMARKS:    _____________________"

    remarks_header.fill = fills["gray"]

    remarks_header.font = fonts["bold"]

    remarks_header.alignment = left_alignment()


    # -----------------------------------
    # TOTALS COLUMNS
    # -----------------------------------

    summary_label_column = (
        total_columns // 2 + 1
    )

    summary_value_column = (
        summary_label_column + 2
    )


    # -----------------------------------
    # SUMMARY LABELS
    # -----------------------------------

    summary_items = [
        "TOTAL NO. OF HEADS:",
        "TOTAL WEIGHT:",
        "AVE WEIGHT:"
    ]

    for index, label in enumerate(
        summary_items
    ):

        row = remarks_row + index


        # -----------------------------------
        # Label
        # -----------------------------------

        label_cell = sheet.cell(
            row=row,
            column=summary_label_column
        )

        label_cell.value = label

        label_cell.fill = fills["gray"]

        label_cell.font = fonts["bold"]

        label_cell.alignment = left_alignment()

        label_cell.border = borders["thin"]


        # -----------------------------------
        # Value
        # -----------------------------------

        value_cell = sheet.cell(
            row=row,
            column=summary_value_column
        )

        value_cell.fill = fills["beige"]

        value_cell.font = fonts["bold"]

        value_cell.alignment = center_alignment()

        value_cell.border = borders["thin"]


    # -----------------------------------
    # WEIGHT DATA RANGES
    # -----------------------------------

    weight_data_ranges = []

    for group_number in range(
        len(groups)
    ):

        weight_column = (
            group_number * 2 + 2
        )

        weight_letter = get_column_letter(
            weight_column
        )

        weight_data_ranges.append(
            f"{weight_letter}{data_start_row}:"
            f"{weight_letter}{data_end_row}"
        )


    # -----------------------------------
    # TOTAL NO. OF HEADS
    # -----------------------------------

    total_heads_formula = (
        "=("
        + "+".join(
            f"COUNT({data_range})"
            for data_range in weight_data_ranges
        )
        + ")*20"
    )

    sheet.cell(
        row=remarks_row,
        column=summary_value_column
    ).value = total_heads_formula


    # -----------------------------------
    # TOTAL WEIGHT
    # -----------------------------------

    sheet.cell(
        row=remarks_row + 1,
        column=summary_value_column
    ).value = (
        "=SUM("
        + ",".join(weight_data_ranges)
        + ")"
    )


    # -----------------------------------
    # AVERAGE WEIGHT
    # -----------------------------------

    sheet.cell(
        row=remarks_row + 2,
        column=summary_value_column
    ).value = (
        f"=IFERROR("
        f"{get_column_letter(summary_value_column)}"
        f"{remarks_row + 1}/"
        f"{get_column_letter(summary_value_column)}"
        f"{remarks_row},"
        f"0)"
    )


    # -----------------------------------
    # Remarks area
    # -----------------------------------

    for row in range(
        remarks_row + 1,
        remarks_row + 3
    ):

        for column in range(
            remarks_start_column,
            remarks_end_column + 1
        ):

            cell = sheet.cell(
                row=row,
                column=column
            )

            cell.border = borders["thin"]


    # -----------------------------------
    # Summary row heights
    # -----------------------------------

    for row in range(
        remarks_row,
        remarks_row + 3
    ):

        sheet.row_dimensions[
            row
        ].height = 24


    # -----------------------------------
    # PLATE / DRIVER SECTION
    # -----------------------------------

    vehicle_header_row = remarks_row + 5

    sheet.merge_cells(
        start_row=vehicle_header_row,
        start_column=1,
        end_row=vehicle_header_row,
        end_column=total_columns
    )

    vehicle_header = sheet.cell(
        row=vehicle_header_row,
        column=1
    )

    vehicle_header.value = "VEHICLE INFORMATION"

    vehicle_header.fill = fills["blue"]

    vehicle_header.font = fonts["summary_header"]

    vehicle_header.alignment = center_alignment()


    # -----------------------------------
    # Plate / Driver
    # -----------------------------------

    vehicle_row = vehicle_header_row + 2

    plate_label_column = 2

    plate_value_column = 3

    driver_label_column = (
        total_columns // 2 + 1
    )

    driver_value_column = (
        driver_label_column + 2
    )


    sheet.cell(
        row=vehicle_row,
        column=plate_label_column
    ).value = "PLATE NO:"

    sheet.cell(
        row=vehicle_row,
        column=plate_label_column
    ).font = fonts["bold"]


    sheet.cell(
        row=vehicle_row,
        column=plate_value_column
    ).value = "____________________"


    sheet.cell(
        row=vehicle_row,
        column=driver_label_column
    ).value = "DRIVER'S NAME:"

    sheet.cell(
        row=vehicle_row,
        column=driver_label_column
    ).font = fonts["bold"]


    sheet.cell(
        row=vehicle_row,
        column=driver_value_column
    ).value = "____________________"


    # -----------------------------------
    # Signature Section
    # -----------------------------------

    signature_header_row = vehicle_row + 4

    sheet.merge_cells(
        start_row=signature_header_row,
        start_column=1,
        end_row=signature_header_row,
        end_column=total_columns
    )

    signature_header = sheet.cell(
        row=signature_header_row,
        column=1
    )

    signature_header.value = "SIGNATURES"

    signature_header.fill = fills["blue"]

    signature_header.font = fonts["summary_header"]

    signature_header.alignment = center_alignment()


    # -----------------------------------
    # Signature positions
    # -----------------------------------

    signature_title_row = (
        signature_header_row + 2
    )

    signature_line_row = (
        signature_title_row + 2
    )

    signature_name_row = (
        signature_line_row + 1
    )


    signature_positions = [
        (
            "RECEIVED BY:",
            1,
            total_columns // 3
        ),
        (
            "GROWER:",
            total_columns // 3 + 1,
            (total_columns // 3) * 2
        ),
        (
            "CHECKED BY:",
            (total_columns // 3) * 2 + 1,
            total_columns
        )
    ]


    for (
        title,
        start_column,
        end_column
    ) in signature_positions:

        # -----------------------------------
        # Signature title
        # -----------------------------------

        sheet.merge_cells(
            start_row=signature_title_row,
            start_column=start_column,
            end_row=signature_title_row,
            end_column=end_column
        )

        title_cell = sheet.cell(
            row=signature_title_row,
            column=start_column
        )

        title_cell.value = title

        title_cell.font = fonts["bold"]

        title_cell.alignment = center_alignment()


        # -----------------------------------
        # Signature line
        # -----------------------------------

        sheet.merge_cells(
            start_row=signature_line_row,
            start_column=start_column,
            end_row=signature_line_row,
            end_column=end_column
        )

        line_cell = sheet.cell(
            row=signature_line_row,
            column=start_column
        )

        line_cell.value = (
            "________________________"
        )

        line_cell.alignment = center_alignment()


        # -----------------------------------
        # Name/signature
        # -----------------------------------

        sheet.merge_cells(
            start_row=signature_name_row,
            start_column=start_column,
            end_row=signature_name_row,
            end_column=end_column
        )

        name_cell = sheet.cell(
            row=signature_name_row,
            column=start_column
        )

        name_cell.value = (
            "(NAME AND SIGNATURE)"
        )

        name_cell.alignment = center_alignment()


    # -----------------------------------
    # Final page row heights
    # -----------------------------------

    sheet.row_dimensions[
        vehicle_header_row
    ].height = 25

    sheet.row_dimensions[
        vehicle_row
    ].height = 30

    sheet.row_dimensions[
        signature_header_row
    ].height = 25

    sheet.row_dimensions[
        signature_title_row
    ].height = 25

    sheet.row_dimensions[
        signature_line_row
    ].height = 30

    sheet.row_dimensions[
        signature_name_row
    ].height = 20