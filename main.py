import re
import glob
from datetime import datetime
from openpyxl import Workbook


# -----------------------------------
# Configuration
# -----------------------------------

output_file = "output.xlsx"


# -----------------------------------
# Find TXT file
# -----------------------------------

txt_files = glob.glob("*.txt")

if not txt_files:
    raise FileNotFoundError("No TXT file found. (VERSION 3)")

if len(txt_files) > 1:
    raise RuntimeError(
        "Multiple TXT files found. Please leave only one TXT file "
        "in the folder."
    )

input_file = txt_files[0]

print(f"Input file: {input_file}")


# -----------------------------------
# 1. Get date from filename
# -----------------------------------

filename_match = re.search(
    r"serial_(\d{8})_\d+\.txt$",
    input_file
)

if not filename_match:
    raise ValueError(
        "Could not determine the date from the filename."
    )

date_text = filename_match.group(1)

file_date = datetime.strptime(
    date_text,
    "%Y%m%d"
).strftime("%m/%d/%Y")

# -----------------------------------
# 2. Read and parse the TXT file
# -----------------------------------

measurements = []
time_started = None

with open(input_file, "r", encoding="utf-8") as file:

    for line in file:

        # Find the Connected time
        connected_match = re.search(
            r"(\d{2}:\d{2}:\d{2}\.\d{3})\s+Connected",
            line
        )

        if connected_match and time_started is None:
            time_started = connected_match.group(1)

        # Find weight measurements
        weight_match = re.search(
            r"(\d{2}:\d{2}:\d{2}\.\d{3})\s+([\d.]+)kg",
            line
        )

        if weight_match:

            time = weight_match.group(1)
            weight = float(weight_match.group(2))

            measurements.append({
                "time": time,
                "weight": weight
            })


# -----------------------------------
# 3. Divide measurements into groups
#    of 10
# -----------------------------------

groups = [
    measurements[i:i + 10]
    for i in range(0, len(measurements), 10)
]


# -----------------------------------
# 4. Create Excel workbook
# -----------------------------------

workbook = Workbook()

sheet = workbook.active
sheet.title = "Weight Data"


# -----------------------------------
# Header
# -----------------------------------

sheet.merge_cells("A1:F1")
sheet["A1"] = "HPS INTEGRATORS CONTRACT ANIMAL GROWING SERVICE"

sheet.merge_cells("A2:F2")
sheet["A2"] = "LIVE SALES TALLY SHEET"


# -----------------------------------
# Information fields
# -----------------------------------

sheet["A4"] = "GROWER:"
sheet["D4"] = "DATE:"
sheet["E4"] = file_date

sheet["A5"] = "LOCATION:"
sheet["D5"] = "BUYER:"

sheet["A6"] = "BUILDING NO:"
sheet["D6"] = "D.R NO:"

sheet["A7"] = "TIME STARTED:"
sheet["B7"] = time_started
sheet["D7"] = "AGE:"

sheet["D8"] = "TIME FINISHED:"


# -----------------------------------
# 6. Create measurement columns
# -----------------------------------

header_row = 10
data_start_row = 11
data_end_row = 20
total_row = 21


for group_number, group in enumerate(groups):

    time_column = group_number * 2 + 1
    weight_column = group_number * 2 + 2


    # -----------------------------------
    # Headers
    # -----------------------------------

    sheet.cell(
        row=header_row,
        column=time_column
    ).value = "TIME"

    sheet.cell(
        row=header_row,
        column=weight_column
    ).value = "WEIGHT"


    # -----------------------------------
    # Measurements
    # -----------------------------------

    for row_number, entry in enumerate(
        group,
        start=data_start_row
    ):

        sheet.cell(
            row=row_number,
            column=time_column
        ).value = entry["time"]

        sheet.cell(
            row=row_number,
            column=weight_column
        ).value = entry["weight"]


    # -----------------------------------
    # Total
    # -----------------------------------

    sheet.cell(
        row=total_row,
        column=time_column
    ).value = "TOTAL"

    sheet.cell(
        row=total_row,
        column=weight_column
    ).value = (
        f"=SUM("
        f"{sheet.cell(row=data_start_row, column=weight_column).coordinate}:"
        f"{sheet.cell(row=data_end_row, column=weight_column).coordinate}"
        f")"
    )


# -----------------------------------
# 7. Column widths
# -----------------------------------

sheet.column_dimensions["A"].width = 15
sheet.column_dimensions["B"].width = 15


# -----------------------------------
# 8. Save Excel
# -----------------------------------

workbook.save(output_file)


# -----------------------------------
# 9. Console output
# -----------------------------------

print(f"Excel file created: {output_file}")
print(f"Date: {file_date.strftime('%m/%d/%Y')}")
print(f"Time started: {time_started}")
print(f"Measurements: {len(measurements)}")
print(f"Groups: {len(groups)}")